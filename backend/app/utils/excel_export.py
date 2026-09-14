import io
from typing import List, Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from fastapi.responses import StreamingResponse


def create_excel_response(
    data: List[Dict[str, Any]],
    sheet_name: str = "Data",
    filename: str = "export.xlsx",
) -> StreamingResponse:
    """Create an Excel file from a list of dicts and return as a streaming response."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    if not data:
        # Empty sheet with header notice
        ws.append(["No data available"])
    else:
        headers = list(data[0].keys())

        # Header style
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # Write headers
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header.replace("_", " ").title())
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # Write data rows
        alt_fill = PatternFill(start_color="EBF2FA", end_color="EBF2FA", fill_type="solid")
        for row_idx, row_data in enumerate(data, 2):
            fill = alt_fill if row_idx % 2 == 0 else None
            for col_idx, key in enumerate(headers, 1):
                value = row_data.get(key, "")
                # Convert None to empty string
                if value is None:
                    value = ""
                cell = ws.cell(row=row_idx, column=col_idx, value=str(value) if not isinstance(value, (int, float, bool)) else value)
                cell.border = thin_border
                if fill:
                    cell.fill = fill

        # Auto-adjust column widths
        for col_idx, header in enumerate(headers, 1):
            col_letter = get_column_letter(col_idx)
            max_length = max(
                len(header) + 2,
                max((len(str(row.get(header, ""))) for row in data), default=0) + 2,
            )
            ws.column_dimensions[col_letter].width = min(max_length, 50)

        # Freeze top row
        ws.freeze_panes = "A2"

    # Save to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
